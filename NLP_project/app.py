from flask import Flask, render_template, request, jsonify
import joblib
import re
import nltk
import io
import zipfile
import json
import csv

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

from nltk.corpus import stopwords
from nltk.stem import WordNetLemmatizer

app = Flask(__name__)

# 50 MB upload limit
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024

nltk.download("stopwords", quiet=True)
nltk.download("wordnet", quiet=True)
nltk.download("omw-1.4", quiet=True)

model = joblib.load("model/news_classifier.pkl")
vectorizer = joblib.load("model/tfidf_vectorizer.pkl")

stop_words = set(stopwords.words("english"))
lemmatizer = WordNetLemmatizer()

SUPPORTED = {
    ".pdf": "PDF",
    ".txt": "Text",
    ".csv": "CSV",
    ".json": "JSON",
    ".docx": "Word",
    ".xlsx": "Excel",
    ".zip": "ZIP archive",
}


def preprocess_text(text):
    text = str(text).lower()
    text = re.sub(r"http\S+|www\S+|https\S+", "", text)
    text = re.sub(r"[^a-zA-Z\s]", "", text)

    words = text.split()
    words = [
        lemmatizer.lemmatize(word)
        for word in words
        if word not in stop_words
    ]
    return " ".join(words)


def extract_pdf(data):
    reader = PdfReader(io.BytesIO(data))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    return text, len(reader.pages)


def extract_docx(data):
    document = Document(io.BytesIO(data))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    # Also read table content
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" ".join(cell.text for cell in row.cells))

    return "\n".join(paragraphs), 1


def extract_xlsx(data):
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []

    for sheet in workbook.worksheets:
        lines.append(f"Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value) for value in row if value is not None]
            if values:
                lines.append(" | ".join(values))

    return "\n".join(lines), len(workbook.worksheets)


def extract_csv(data):
    text = data.decode("utf-8-sig", errors="replace")

    try:
        rows = csv.reader(io.StringIO(text))
        lines = []
        for row in rows:
            lines.append(" ".join(str(cell) for cell in row))
        return "\n".join(lines), 1
    except Exception:
        return text, 1


def extract_json(data):
    text = data.decode("utf-8-sig", errors="replace")

    try:
        obj = json.loads(text)
        return json.dumps(obj, ensure_ascii=False, indent=2), 1
    except Exception:
        return text, 1


def safe_zip_member(name):
    # Prevent zip path traversal.
    normalized = name.replace("\\", "/")
    return (
        not normalized.startswith("/")
        and ".." not in normalized.split("/")
    )


def extract_zip(data):
    all_text = []
    file_count = 0
    ignored = []

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for info in archive.infolist():

            if info.is_dir():
                continue

            if not safe_zip_member(info.filename):
                ignored.append(info.filename)
                continue

            suffix = "." + info.filename.rsplit(".", 1)[-1].lower() if "." in info.filename else ""

            if suffix not in SUPPORTED or suffix == ".zip":
                ignored.append(info.filename)
                continue

            try:
                member_data = archive.read(info)
                text, _ = extract_by_extension(suffix, member_data)

                if text.strip():
                    all_text.append(
                        f"\n--- {info.filename} ---\n{text}"
                    )
                    file_count += 1

            except Exception:
                ignored.append(info.filename)

    return "\n".join(all_text), file_count, ignored


def extract_by_extension(suffix, data):
    if suffix == ".pdf":
        return extract_pdf(data)

    if suffix == ".txt":
        return data.decode("utf-8-sig", errors="replace"), 1

    if suffix == ".csv":
        return extract_csv(data)

    if suffix == ".json":
        return extract_json(data)

    if suffix == ".docx":
        return extract_docx(data)

    if suffix == ".xlsx":
        return extract_xlsx(data)

    if suffix == ".zip":
        text, count, _ = extract_zip(data)
        return text, count

    raise ValueError(f"Unsupported file type: {suffix}")


def analyze_article(article):
    clean_article = preprocess_text(article)
    article_vector = vectorizer.transform([clean_article])

    prediction = model.predict(article_vector)[0]
    probabilities = model.predict_proba(article_vector)[0]

    classes = list(model.classes_)
    probability_data = [
        {
            "category": str(category),
            "probability": round(float(prob) * 100, 2)
        }
        for category, prob in zip(classes, probabilities)
    ]
    probability_data.sort(key=lambda x: x["probability"], reverse=True)

    feature_names = vectorizer.get_feature_names_out()
    row = article_vector.toarray()[0]
    top_indices = row.argsort()[::-1]

    keywords = []
    for idx in top_indices:
        if row[idx] <= 0:
            break
        term = feature_names[idx]
        if term not in keywords:
            keywords.append(term)
        if len(keywords) == 8:
            break

    words = article.split()
    word_count = len(words)
    char_count = len(article)
    reading_time = max(1, round(word_count / 200))
    confidence = max(probabilities) * 100

    if confidence >= 80:
        confidence_level = "High confidence"
    elif confidence >= 60:
        confidence_level = "Moderate confidence"
    else:
        confidence_level = "Low confidence"

    return {
        "category": str(prediction),
        "confidence": round(float(confidence), 2),
        "confidence_level": confidence_level,
        "probabilities": probability_data,
        "keywords": keywords,
        "word_count": word_count,
        "char_count": char_count,
        "reading_time": reading_time,
    }


@app.route("/")
def home():
    return render_template("index.html")


@app.route("/predict-file", methods=["POST"])
def predict_file():
    if "file" not in request.files:
        return jsonify({"error": "Please select a file."}), 400

    uploaded = request.files["file"]

    if not uploaded.filename:
        return jsonify({"error": "Please select a file."}), 400

    filename = uploaded.filename
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if suffix not in SUPPORTED:
        return jsonify({
            "error": (
                "Unsupported file type. Supported: "
                "PDF, TXT, CSV, JSON, DOCX, XLSX and ZIP."
            )
        }), 400

    try:
        raw_data = uploaded.read()

        if suffix == ".zip":
            article, item_count, ignored = extract_zip(raw_data)
            pages = item_count
        else:
            article, pages = extract_by_extension(suffix, raw_data)
            item_count = 1
            ignored = []

        article = article.strip()

        if not article:
            return jsonify({
                "error": (
                    "No readable text was found in this file. "
                    "Scanned/image-only PDFs require OCR."
                )
            }), 400

        result = analyze_article(article)
        result.update({
            "filename": filename,
            "file_type": SUPPORTED[suffix],
            "pages": pages,
            "items": item_count,
            "extracted_preview": article[:3500],
            "ignored_files": ignored[:10],
        })

        return jsonify(result)

    except Exception as e:
        return jsonify({
            "error": f"Could not process the file: {str(e)}"
        }), 500


@app.errorhandler(413)
def too_large(error):
    return jsonify({
        "error": "File is too large. Maximum size is 50 MB."
    }), 413


if __name__ == "__main__":
    app.run(debug=True)
